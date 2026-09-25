# One-off script: builds the 2-page sponsorship proposal as a Word doc.
# Run: python scripts/build_sponsor_proposal.py
# Kept in the repo so it can be regenerated whenever the tiers, contact
# info, or site link change — edit the CONTENT section below and rerun.
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE

# ---- brand palette (matches the site's tokens.css, light theme) ----
CRIMSON = RGBColor(0xC2, 0x21, 0x3A)
GOLD    = RGBColor(0xD9, 0x70, 0x3C)
INK     = RGBColor(0x23, 0x16, 0x1A)
INK_SOFT = RGBColor(0x5C, 0x41, 0x48)
CREAM   = "FBF6F1"
LINE    = "E7D8D2"

SITE_URL = "https://roosewelt-nacth.github.io/Vimusement/sponsors.html"
SITE_NOTE = "the site is on a temporary/testing link right now — swap this before final print"

def add_hyperlink(paragraph, url, text, color=CRIMSON, underline=True):
    part = paragraph.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    if color:
        c = OxmlElement("w:color")
        c.set(qn("w:val"), "%02X%02X%02X" % (color[0], color[1], color[2]))
        rPr.append(c)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rPr.append(u)
    new_run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

def shade_paragraph(paragraph, fill_hex):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    pPr.append(shd)

def left_border(paragraph, color_hex, size=24):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), color_hex)
    pBdr.append(left)
    pPr.append(pBdr)

def bottom_border(paragraph, color_hex, size=8, space=8):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color_hex)
    pBdr.append(bottom)
    pPr.append(pBdr)

def set_run(run, size=10.5, color=INK, bold=False, italic=False, font="Calibri", caps=False):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font
    run.font.all_caps = caps

def para(doc, space_before=0, space_after=8, indent=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if indent is not None:
        p.paragraph_format.left_indent = Cm(indent)
    return p

def eyebrow(doc):
    p = para(doc, space_after=18)
    r = p.add_run("AN INITIATIVE OF VICTORIANS YOUTH")
    set_run(r, size=8.5, color=GOLD, bold=True, caps=True)

def title(doc, text, size=30):
    p = para(doc, space_after=2)
    r = p.add_run(text)
    set_run(r, size=size, color=INK, bold=True, font="Georgia")

def subtitle(doc, text):
    p = para(doc, space_after=14)
    r = p.add_run(text)
    set_run(r, size=12.5, color=INK_SOFT)

def rule(doc):
    p = para(doc, space_after=14)
    bottom_border(p, "C2213A", size=16, space=6)

def detail_row(doc, label, text):
    p = para(doc, space_after=10)
    r1 = p.add_run(label + "\t")
    set_run(r1, size=9, color=GOLD, bold=True, caps=True)
    r2 = p.add_run(text)
    set_run(r2, size=10.3, color=INK)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(2.7))
    p.paragraph_format.left_indent = Cm(2.7)
    p.paragraph_format.first_line_indent = Cm(-2.7)

def bullet(doc, text):
    p = para(doc, space_after=6)
    p.paragraph_format.left_indent = Cm(0.6)
    r = p.add_run("•   " + text)
    set_run(r, size=10.3, color=INK)

def quote_block(doc, text):
    p = para(doc, space_before=6, space_after=6, indent=0.4)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(10)
    shade_paragraph(p, CREAM)
    left_border(p, "D9703C", size=24)
    r = p.add_run("“" + text + "”")
    set_run(r, size=11, color=INK, italic=True, font="Georgia")

def tier_block(doc, name, amount, benefit, last=False):
    p = para(doc, space_before=10, space_after=2)
    r1 = p.add_run(name + "   ")
    set_run(r1, size=12.5, color=INK, bold=True, font="Georgia")
    r2 = p.add_run(amount)
    set_run(r2, size=10.5, color=CRIMSON, bold=True, italic=True)
    p2 = para(doc, space_after=12 if last else 16)
    r3 = p2.add_run(benefit)
    set_run(r3, size=9.8, color=INK_SOFT)
    if not last:
        bottom_border(p2, LINE, size=6, space=10)

# ============================================================
doc = Document()
section = doc.sections[0]
section.page_width = Cm(21.0)
section.page_height = Cm(29.7)
for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
    setattr(section, m, Cm(2.0))
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(10.3)

# ---- PAGE 1 — the event ----
eyebrow(doc)
title(doc, "VIMUSEMENT 2026")
subtitle(doc, "A Community Fundraising Carnival")
rule(doc)

detail_row(doc, "Where", "Ascension Church, Metha Nagar, Aminjikkarai, Chennai")
detail_row(doc, "Who", "Families and young people from right across the neighbourhood — hundreds through the gates in one day")
detail_row(doc, "Purpose", "Scholarships, a medical-emergency fund, and hardship support for neighbours in need — 100% of what's raised, after costs, goes to the cause")

p = para(doc, space_before=10, space_after=8)
r = p.add_run("Why partner with us")
set_run(r, size=15, color=INK, bold=True, font="Georgia")

for line in [
    "A local, built-in crowd — no footfall to chase, just show up and be seen",
    "A family + youth audience, right around Aminjikkarai",
    "Real on-ground brand visibility: posters, banners, stalls",
    "Social media exposure across our event pages",
    "MC and stage announcements naming your brand",
    "Product sampling or distribution to the crowd",
    "Brand activation space, if you want to bring one",
]:
    bullet(doc, line)

doc.add_paragraph().paragraph_format.space_after = Pt(6)
quote_block(doc, "We're not just asking for funding — we can give your brand real visibility "
                  "through our posters, social media, stage announcements and, depending on the "
                  "partnership, product integration.")

p = para(doc, space_before=16)
r = p.add_run("Partnership options overleaf →")
set_run(r, size=10, color=INK_SOFT)

doc.add_page_break()

# ============================================================
# ---- PAGE 2 — partnership options ----
eyebrow(doc)
title(doc, "Partnership options", size=24)
subtitle(doc, "Pick what fits, or tell us what works for your business")
rule(doc)

tier_block(doc, "Community Partner", "Flexible",
    "Logo on event banners, a mention in our social media posts, and a shout-out "
    "on stage — let's talk about what works for you.")
tier_block(doc, "Product Partner", "Products / vouchers",
    "Give what you make or sell — food, drinks, goods — and we promote your "
    "brand in return. Great for restaurants, cafés and shops.")
tier_block(doc, "Prize Partner", "Gifts / vouchers",
    "Sponsor a game prize — your brand gets named the moment it's won.", last=True)

p = para(doc, space_before=4, space_after=16)
r = p.add_run("Not sure which fits? Message us and we'll shape a package around what "
              "works for your business — a mix of the above is always an option.")
set_run(r, size=10, color=INK_SOFT)

# contact box
p = para(doc, space_before=8, space_after=2)
shade_paragraph(p, CREAM)
left_border(p, "C2213A", size=24)
r = p.add_run("GET IN TOUCH")
set_run(r, size=8.5, color=GOLD, bold=True, caps=True)

p = para(doc, space_after=2)
shade_paragraph(p, CREAM)
left_border(p, "C2213A", size=24)
r = p.add_run("Austin Prince Roosewelt")
set_run(r, size=13, color=INK, bold=True)
r = p.add_run("  ·  Victorians Youth")
set_run(r, size=10.3, color=INK_SOFT)

p = para(doc, space_after=2)
shade_paragraph(p, CREAM)
left_border(p, "C2213A", size=24)
r = p.add_run("WhatsApp: +91 63794 68686")
set_run(r, size=10.3, color=INK)

p = para(doc, space_after=10)
shade_paragraph(p, CREAM)
left_border(p, "C2213A", size=24)
r = p.add_run("Web: ")
set_run(r, size=10.3, color=INK)
add_hyperlink(p, SITE_URL, SITE_URL, color=CRIMSON)
p.add_run().add_break()

rule2 = para(doc, space_before=14, space_after=8)
bottom_border(rule2, LINE, size=6, space=0)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Vimusement 2026 · An annual fundraiser by the parish community, "
              "Ascension Church, Aminjikkarai.")
set_run(r, size=8, color=INK_SOFT)

out_path = "assets/docs/Vimusement-2026-Sponsorship-Proposal.docx"
doc.save(out_path)
print("saved", out_path)
